# from flask import Flask, request, jsonify
# from flask_cors import CORS

# import os
# import re
# import docx
# import fitz  # PyMuPDF
# from PIL import Image
# import pytesseract
# from collections import defaultdict


# import spacy
# from spacy.matcher import PhraseMatcher
# from rapidfuzz import process, fuzz


# app = Flask(__name__)
# CORS(app)
# # CORS(app, resources={r"/analyze": {"origins": "http://localhost:5173"}})

# # ---- Initialize spaCy model ----
# nlp = spacy.load("en_core_web_sm")

# # ---- Skill pool with some expansions ----
# SKILL_POOL = {
#     "Software & Programming": {
#         "score": 5,
#         "keywords": [
#             "python", "java", "c++", "c", "c#", "javascript", "sql", "git", "rest api",
#             "ci/cd", "agile", "scrum", "jira", "nodejs", "express", "flutter"
#         ]
#     },
#     "Data Science & AI": {
#         "score": 10,
#         "keywords": [
#             "machine learning", "deep learning", "nlp", "tensorflow", "pytorch", "scikit-learn",
#             "pandas", "numpy", "matplotlib", "seaborn", "r", "spark", "hadoop", "bi",
#             "tableau", "power bi"
#         ]
#     },
#     "Electrical & Electronics Engineering": {
#         "score": 10,
#         "keywords": [
#             "verilog", "embedded c", "matlab", "xilinx vivado", "altium designer",
#             "eagle", "pcb design", "circuit design", "vsim", "arduino", "raspberry pi",
#             "microcontrollers", "vlsi", "fpga", "hfss", "ece"
#         ]
#     },
#     "Mechanical & Aerospace Engineering": {
#         "score": 8,
#         "keywords": [
#             "autocad", "solidworks", "ansys", "catia", "cad", "cam", "thermodynamics",
#             "fluid mechanics", "materials science", "fea", "robotics",
#             "3d printing", "propulsion", "aerodynamics"
#         ]
#     },
#     "Civil Engineering": {
#         "score": 8,
#         "keywords": [
#             "autocad civil 3d", "revit", "etabs", "staad pro", "safe", "gis",
#             "structural analysis", "geotechnical engineering", "transportation engineering"
#         ]
#     },
#     "Web Development": {
#         "score": 5,
#         "keywords": [
#             "html", "css", "javascript", "react", "angular", "vue.js", "node.js", "django", "flask"
#         ]
#     },
#     "Project Management & Business": {
#         "score": 4,
#         "keywords": [
#             "project management", "scrum", "agile", "jira", "confluence", "microsoft office",
#             "excel", "powerpoint", "finance", "marketing"
#         ]
#     }
# }
# FLATTENED_SKILL_POOL = [skill for sublist in [d['keywords'] for d in SKILL_POOL.values()] for skill in sublist]

# # ---- Section keywords expanded ----
# SECTION_KEYWORDS = {
#     "experience": [
#         "experience", "work experience", "employment history", "professional experience",
#         "work history", "roles", "career", "professional background"
#     ],
#     "education": [
#         "education", "academic", "academic background", "qualifications", "degrees",
#         "education & qualifications", "schooling", "academics"
#     ],
#     "skills": [
#         "skills", "technical skills", "key skills", "competencies", "expertise", "skillset"
#     ],
#     "projects": [
#         "projects", "personal projects", "academic projects", "selected projects", "research"
#     ],
#     "certifications": ["certifications", "licenses"],
#     "summary": [
#         "summary", "profile", "professional summary", "about", "about me", "career objective"
#     ],
#     "awards": ["awards", "honors", "achievements"],
#     "publications": ["publications", "research papers"]
# }

# phrase_matcher = PhraseMatcher(nlp.vocab, attr="LOWER")
# for sec, kws in SECTION_KEYWORDS.items():
#     phrase_matcher.add(sec, [nlp.make_doc(k) for k in kws])

# RE_EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", re.I)
# RE_DATE_RANGE = re.compile(
#     r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\s+\d{4}\s*[\-–]\s*(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\s+\d{4}|(present|current|until now)", re.I
# )

# # -------- Utility Functions --------
# def normalize_text(text):
#     if not text:
#         return ""
#     text = (
#         text.replace("\u2013", "-")
#         .replace("\u2014", "-")
#         .replace("\u2018", "'")
#         .replace("\u2019", "'")
#         .replace("\u201c", '"')
#         .replace("\u201d", '"')
#     )
#     text = re.sub(r"\r\n?", "\n", text)
#     text = re.sub(r"[ \t]+", " ", text)
#     lines = [ln.strip() for ln in text.splitlines()]
#     return "\n".join(lines).strip()

# def extract_text_from_docx(path):
#     try:
#         doc = docx.Document(path)
#         full_text = []

#         # Paragraphs
#         for para in doc.paragraphs:
#             full_text.append(para.text)

#         # Tables - concatenated with '|' per row
#         for table in doc.tables:
#             for row in table.rows:
#                 row_text = []
#                 for cell in row.cells:
#                     row_text.append(cell.text)
#                 full_text.append(" | ".join(row_text))

#         return "\n".join(full_text)
#     except Exception:
#         return ""

# def extract_text_from_pdf(path):
#     try:
#         doc = fitz.open(path)
#         full_text = []
#         for page in doc:
#             text = page.get_text("text") or ""
#             if text.strip():
#                 full_text.append(text)
#             else:
#                 # Use OCR fallback if page text empty
#                 pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
#                 img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
#                 ocr_text = pytesseract.image_to_string(img)
#                 full_text.append(ocr_text)
#         return "\n".join(full_text)
#     except Exception:
#         return ""

# def fuzzy_section_heading(line, threshold=80):
#     if not line:
#         return None
#     doc_line = nlp(line)
#     matches = phrase_matcher(doc_line)
#     if matches:
#         match_id, start, end = matches[0]
#         return nlp.vocab.strings[match_id]
#     best = None
#     best_score = 0
#     for sec, kws in SECTION_KEYWORDS.items():
#         for kw in kws:
#             score = fuzz.token_set_ratio(line.lower(), kw.lower())
#             if score > best_score:
#                 best_score = score
#                 best = sec
#     if best_score >= threshold:
#         return best
#     return None

# def fuzzy_find_skills(text, pool=FLATTENED_SKILL_POOL, limit=30, threshold=80):
#     found = set()
#     doc = nlp(text.lower())
#     skill_matcher = PhraseMatcher(nlp.vocab, attr="LOWER")
#     skill_patterns = [nlp.make_doc(s) for s in pool]
#     skill_matcher.add("SKILLS", skill_patterns)
#     matches = skill_matcher(doc)
#     for match_id, start, end in matches:
#         found.add(doc[start:end].text)
#     if len(found) < 10:
#         tokens = [token.text for token in doc if not token.is_punct and not token.is_space and len(token) > 2]
#         for token in tokens:
#             best = process.extractOne(token, pool, scorer=fuzz.token_set_ratio)
#             if best and best[1] >= threshold:
#                 found.add(best[0].lower())
#     return sorted(list(found))[:limit]

# def get_specialization(skills, education_lines, experience_lines):
#     skill_set = set(skills)
#     scores = defaultdict(int)
#     for category, data in SKILL_POOL.items():
#         for keyword in data["keywords"]:
#             if keyword in skill_set:
#                 scores[category] += data["score"]
#     education_text = " ".join(education_lines).lower()
#     experience_text = " ".join(experience_lines).lower()
#     if any(s in education_text for s in ["computer science", "cs", "software"]):
#         scores["Software & Programming"] += 15
#     if any(s in education_text for s in ["electronics", "ece"]):
#         scores["Electrical & Electronics Engineering"] += 15
#     if any(s in education_text for s in ["mechanical", "aerospace"]):
#         scores["Mechanical & Aerospace Engineering"] += 15
#     if any(s in education_text for s in ["civil engineering", "civil"]):
#         scores["Civil Engineering"] += 15
#     if any(s in experience_text for s in ["data scientist", "machine learning"]):
#         scores["Data Science & AI"] += 15
#     if any(s in experience_text for s in ["web developer", "full-stack"]):
#         scores["Web Development"] += 10
#     total_score = sum(scores.values())
#     if not total_score:
#         return "General Professional", {}
#     weighted_scores = {k: (v / total_score) * 100 for k, v in scores.items()}
#     best_category = max(weighted_scores, key=weighted_scores.get)
#     mapping = {
#         "Software & Programming": "Software Engineer / Developer",
#         "Data Science & AI": "Machine Learning Engineer / Data Scientist",
#         "Electrical & Electronics Engineering": "Electronics Engineer / ECE Professional",
#         "Mechanical & Aerospace Engineering": "Mechanical Engineer",
#         "Civil Engineering": "Civil Engineer",
#         "Web Development": "Web Developer",
#         "Project Management & Business": "Project Management / Business Analyst",
#     }
#     return mapping.get(best_category, "General Professional"), weighted_scores

# def parse_resume_text(raw_text):
#     text = normalize_text(raw_text)
#     lines = [ln for ln in text.splitlines() if ln.strip()]

#     def extract_name(lines):
#         for ln in lines[:15]:
#             ln = ln.strip()
#             match = re.match(r"^([A-Z][a-z]+(?:\s[A-Z][a-z]+)+)$", ln)
#             if match:
#                 return match.group(1).strip()
#             if ln.isupper() and len(ln.split()) >= 2 and all(w.isalpha() for w in ln.split()):
#                 return ln.title()
#         for ln in lines[:20]:
#             if "name" in ln.lower():
#                 parts = ln.split(":")
#                 if len(parts) > 1:
#                     return parts[1].strip().title()
#         return "Not Found"

#     def extract_emails(text):
#         return re.findall(RE_EMAIL, text) or ["Not Found"]

#     def extract_phones(text):
#         phone_pattern = r"(\+?\d{1,3}[-\s]?)?\(?\d{2,4}\)?[-\s]?\d{3,4}[-\s]?\d{4}"
#         matches = re.findall(phone_pattern, text)
#         results = []
#         for m in matches:
#             num = "".join(m) if isinstance(m, tuple) else m
#             digits = re.sub(r"\D", "", num)
#             if 8 <= len(digits) <= 15:
#                 results.append(num.strip())
#         return list(dict.fromkeys(results)) or ["Not Found"]

#     name = extract_name(lines)
#     emails = extract_emails(raw_text)
#     phones = extract_phones(raw_text)

#     sections = defaultdict(list)
#     current_section = "preamble"
#     for ln in lines:
#         heading = fuzzy_section_heading(ln)
#         if heading:
#             current_section = heading
#             continue
#         sections[current_section].append(ln)
#     section_text = {k: "\n".join(v).strip() for k, v in sections.items()}

#     edu_lines = [ln for ln in section_text.get("education", "").splitlines() if len(ln) > 3]
#     if not edu_lines:
#         for ln in lines:
#             if any(
#                 d.lower() in ln.lower()
#                 for d in [
#                     "bachelor",
#                     "master",
#                     "phd",
#                     "degree",
#                     "university",
#                     "college",
#                     "institute",
#                     "school",
#                 ]
#             ):
#                 edu_lines.append(ln)
#     education = list(dict.fromkeys(edu_lines))[:8]

#     experience_lines = [ln for ln in section_text.get("experience", "").splitlines() if len(ln) > 3]
#     if not experience_lines:
#         for ln in lines:
#             if any(
#                 k in ln.lower()
#                 for k in ["engineer", "developer", "analyst", "intern", "project", "worked"]
#             ) or re.search(RE_DATE_RANGE, ln):
#                 experience_lines.append(ln)
#     experience = list(dict.fromkeys(experience_lines))[:15]

#     skills = fuzzy_find_skills(raw_text, FLATTENED_SKILL_POOL, threshold=78)
#     specialization, scores = get_specialization(skills, education, experience)

#     summary_parts = []
#     if name != "Not Found":
#         summary_parts.append(f"A profile for {name} has been successfully parsed. The system uses an advanced NLP pipeline.")
#     if education:
#         summary_parts.append(f"Education highlights include: {education[0] + (', ...' if len(education) > 1 else '')}.")
#     if experience:
#         summary_parts.append(f"The experience section details roles such as '{experience[0].split('(')[0].strip()}' and mentions key projects.")
#     if skills:
#         summary_parts.append(
#             f"The candidate's core skills are in **{specialization}**, with a strong command of: {', '.join(skills[:6]) + '...' if len(skills) > 6 else ''}."
#         )
#     summary_parts.append(f"Overall, the resume suggests a strong fit for a **{specialization}** role.")

#     summary = "\n\n".join(summary_parts)

#     return {
#         "Name": name,
#         "Emails": emails,
#         "Phones": phones,
#         "Education": education,
#         "Experience": experience,
#         "Skills": skills,
#         "Specialization": specialization,
#         "SpecializationScores": scores,
#         "Summary": summary,
#     }

# # -------- Flask App --------
# app = Flask(__name__)

# @app.route('/analyze', methods=['POST'])
# def analyze_resume():
#     if 'file' not in request.files:
#         return jsonify({"error": "No file uploaded."}), 400
#     file = request.files['file']
#     ext = os.path.splitext(file.filename)[1].lower()
#     temp_path = "temp" + ext
#     file.save(temp_path)
#     try:
#         if ext == ".pdf":
#             text = extract_text_from_pdf(temp_path)
#         elif ext == ".docx":
#             text = extract_text_from_docx(temp_path)
#         elif ext == ".txt":
#             with open(temp_path, "r", encoding="utf-8", errors="ignore") as f:
#                 text = f.read()
#         else:
#             os.remove(temp_path)
#             return jsonify({"error": "Unsupported file type"}), 400
#         result = parse_resume_text(text)
#         os.remove(temp_path)
#         return jsonify(result)
#     except Exception as e:
#         os.remove(temp_path)
#         return jsonify({"error": str(e)}), 500

# if __name__ == "__main__":
#     app.run(debug=True, port=5001)










from flask import Flask, request, jsonify
from flask_cors import CORS

import os
import re
import docx
import fitz  # PyMuPDF
from PIL import Image
import pytesseract
from collections import defaultdict

import spacy
from spacy.matcher import PhraseMatcher
from rapidfuzz import process, fuzz


# -------- Flask App --------
app = Flask(__name__)
CORS(app, resources={r"/*": {"origins": "http://localhost:5173"}})

# ---- Initialize spaCy model ----
nlp = spacy.load("en_core_web_sm")

# ---- Skill pool with some expansions ----
SKILL_POOL = {
    "Software & Programming": {
        "score": 5,
        "keywords": [
            "python", "java", "c++", "c", "c#", "javascript", "sql", "git", "rest api",
            "ci/cd", "agile", "scrum", "jira", "nodejs", "express", "flutter"
        ]
    },
    "Data Science & AI": {
        "score": 10,
        "keywords": [
            "machine learning", "deep learning", "nlp", "tensorflow", "pytorch", "scikit-learn",
            "pandas", "numpy", "matplotlib", "seaborn", "r", "spark", "hadoop", "bi",
            "tableau", "power bi"
        ]
    },
    "Electrical & Electronics Engineering": {
        "score": 10,
        "keywords": [
            "verilog", "embedded c", "matlab", "xilinx vivado", "altium designer",
            "eagle", "pcb design", "circuit design", "vsim", "arduino", "raspberry pi",
            "microcontrollers", "vlsi", "fpga", "hfss", "ece"
        ]
    },
    "Mechanical & Aerospace Engineering": {
        "score": 8,
        "keywords": [
            "autocad", "solidworks", "ansys", "catia", "cad", "cam", "thermodynamics",
            "fluid mechanics", "materials science", "fea", "robotics",
            "3d printing", "propulsion", "aerodynamics"
        ]
    },
    "Civil Engineering": {
        "score": 8,
        "keywords": [
            "autocad civil 3d", "revit", "etabs", "staad pro", "safe", "gis",
            "structural analysis", "geotechnical engineering", "transportation engineering"
        ]
    },
    "Web Development": {
        "score": 5,
        "keywords": [
            "html", "css", "javascript", "react", "angular", "vue.js", "node.js", "django", "flask"
        ]
    },
    "Project Management & Business": {
        "score": 4,
        "keywords": [
            "project management", "scrum", "agile", "jira", "confluence", "microsoft office",
            "excel", "powerpoint", "finance", "marketing"
        ]
    }
}
FLATTENED_SKILL_POOL = [skill for sublist in [d['keywords'] for d in SKILL_POOL.values()] for skill in sublist]

# ---- Section keywords expanded ----
SECTION_KEYWORDS = {
    "experience": [
        "experience", "work experience", "employment history", "professional experience",
        "work history", "roles", "career", "professional background"
    ],
    "education": [
        "education", "academic", "academic background", "qualifications", "degrees",
        "education & qualifications", "schooling", "academics"
    ],
    "skills": [
        "skills", "technical skills", "key skills", "competencies", "expertise", "skillset"
    ],
    "projects": [
        "projects", "personal projects", "academic projects", "selected projects", "research"
    ],
    "certifications": ["certifications", "licenses"],
    "summary": [
        "summary", "profile", "professional summary", "about", "about me", "career objective"
    ],
    "awards": ["awards", "honors", "achievements"],
    "publications": ["publications", "research papers"]
}

phrase_matcher = PhraseMatcher(nlp.vocab, attr="LOWER")
for sec, kws in SECTION_KEYWORDS.items():
    phrase_matcher.add(sec, [nlp.make_doc(k) for k in kws])

RE_EMAIL = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}", re.I)
RE_DATE_RANGE = re.compile(
    r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\s+\d{4}\s*[\-–]\s*(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)\s+\d{4}|(present|current|until now)", re.I
)

# -------- Utility Functions --------
def normalize_text(text):
    if not text:
        return ""
    text = (
        text.replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2018", "'")
        .replace("\u2019", "'")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
    )
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    lines = [ln.strip() for ln in text.splitlines()]
    return "\n".join(lines).strip()

def extract_text_from_docx(path):
    try:
        doc = docx.Document(path)
        full_text = []

        # Paragraphs
        for para in doc.paragraphs:
            full_text.append(para.text)

        # Tables - concatenated with '|' per row
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(cell.text)
                full_text.append(" | ".join(row_text))

        return "\n".join(full_text)
    except Exception:
        return ""

def extract_text_from_pdf(path):
    try:
        doc = fitz.open(path)
        full_text = []
        for page in doc:
            text = page.get_text("text") or ""
            if text.strip():
                full_text.append(text)
            else:
                # Use OCR fallback if page text empty
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_text = pytesseract.image_to_string(img)
                full_text.append(ocr_text)
        return "\n".join(full_text)
    except Exception:
        return ""

def fuzzy_section_heading(line, threshold=80):
    if not line:
        return None
    doc_line = nlp(line)
    matches = phrase_matcher(doc_line)
    if matches:
        match_id, start, end = matches[0]
        return nlp.vocab.strings[match_id]
    best = None
    best_score = 0
    for sec, kws in SECTION_KEYWORDS.items():
        for kw in kws:
            score = fuzz.token_set_ratio(line.lower(), kw.lower())
            if score > best_score:
                best_score = score
                best = sec
    if best_score >= threshold:
        return best
    return None

def fuzzy_find_skills(text, pool=FLATTENED_SKILL_POOL, limit=30, threshold=80):
    found = set()
    doc = nlp(text.lower())
    skill_matcher = PhraseMatcher(nlp.vocab, attr="LOWER")
    skill_patterns = [nlp.make_doc(s) for s in pool]
    skill_matcher.add("SKILLS", skill_patterns)
    matches = skill_matcher(doc)
    for match_id, start, end in matches:
        found.add(doc[start:end].text)
    if len(found) < 10:
        tokens = [token.text for token in doc if not token.is_punct and not token.is_space and len(token) > 2]
        for token in tokens:
            best = process.extractOne(token, pool, scorer=fuzz.token_set_ratio)
            if best and best[1] >= threshold:
                found.add(best[0].lower())
    return sorted(list(found))[:limit]

def get_specialization(skills, education_lines, experience_lines):
    skill_set = set(skills)
    scores = defaultdict(int)
    for category, data in SKILL_POOL.items():
        for keyword in data["keywords"]:
            if keyword in skill_set:
                scores[category] += data["score"]
    education_text = " ".join(education_lines).lower()
    experience_text = " ".join(experience_lines).lower()
    if any(s in education_text for s in ["computer science", "cs", "software"]):
        scores["Software & Programming"] += 15
    if any(s in education_text for s in ["electronics", "ece"]):
        scores["Electrical & Electronics Engineering"] += 15
    if any(s in education_text for s in ["mechanical", "aerospace"]):
        scores["Mechanical & Aerospace Engineering"] += 15
    if any(s in education_text for s in ["civil engineering", "civil"]):
        scores["Civil Engineering"] += 15
    if any(s in experience_text for s in ["data scientist", "machine learning"]):
        scores["Data Science & AI"] += 15
    if any(s in experience_text for s in ["web developer", "full-stack"]):
        scores["Web Development"] += 10
    total_score = sum(scores.values())
    if not total_score:
        return "General Professional", {}
    weighted_scores = {k: (v / total_score) * 100 for k, v in scores.items()}
    best_category = max(weighted_scores, key=weighted_scores.get)
    mapping = {
        "Software & Programming": "Software Engineer / Developer",
        "Data Science & AI": "Machine Learning Engineer / Data Scientist",
        "Electrical & Electronics Engineering": "Electronics Engineer / ECE Professional",
        "Mechanical & Aerospace Engineering": "Mechanical Engineer",
        "Civil Engineering": "Civil Engineer",
        "Web Development": "Web Developer",
        "Project Management & Business": "Project Management / Business Analyst",
    }
    return mapping.get(best_category, "General Professional"), weighted_scores

def parse_resume_text(raw_text):
    text = normalize_text(raw_text)
    lines = [ln for ln in text.splitlines() if ln.strip()]

    def extract_name(lines):
        for ln in lines[:15]:
            ln = ln.strip()
            match = re.match(r"^([A-Z][a-z]+(?:\s[A-Z][a-z]+)+)$", ln)
            if match:
                return match.group(1).strip()
            if ln.isupper() and len(ln.split()) >= 2 and all(w.isalpha() for w in ln.split()):
                return ln.title()
        for ln in lines[:20]:
            if "name" in ln.lower():
                parts = ln.split(":")
                if len(parts) > 1:
                    return parts[1].strip().title()
        return "Not Found"

    def extract_emails(text):
        return re.findall(RE_EMAIL, text) or ["Not Found"]

    def extract_phones(text):
        phone_pattern = r"(\+?\d{1,3}[-\s]?)?\(?\d{2,4}\)?[-\s]?\d{3,4}[-\s]?\d{4}"
        matches = re.findall(phone_pattern, text)
        results = []
        for m in matches:
            num = "".join(m) if isinstance(m, tuple) else m
            digits = re.sub(r"\D", "", num)
            if 8 <= len(digits) <= 15:
                results.append(num.strip())
        return list(dict.fromkeys(results)) or ["Not Found"]

    name = extract_name(lines)
    emails = extract_emails(raw_text)
    phones = extract_phones(raw_text)

    sections = defaultdict(list)
    current_section = "preamble"
    for ln in lines:
        heading = fuzzy_section_heading(ln)
        if heading:
            current_section = heading
            continue
        sections[current_section].append(ln)
    section_text = {k: "\n".join(v).strip() for k, v in sections.items()}

    edu_lines = [ln for ln in section_text.get("education", "").splitlines() if len(ln) > 3]
    if not edu_lines:
        for ln in lines:
            if any(
                d.lower() in ln.lower()
                for d in [
                    "bachelor",
                    "master",
                    "phd",
                    "degree",
                    "university",
                    "college",
                    "institute",
                    "school",
                ]
            ):
                edu_lines.append(ln)
    education = list(dict.fromkeys(edu_lines))[:8]

    experience_lines = [ln for ln in section_text.get("experience", "").splitlines() if len(ln) > 3]
    if not experience_lines:
        for ln in lines:
            if any(
                k in ln.lower()
                for k in ["engineer", "developer", "analyst", "intern", "project", "worked"]
            ) or re.search(RE_DATE_RANGE, ln):
                experience_lines.append(ln)
    experience = list(dict.fromkeys(experience_lines))[:15]

    skills = fuzzy_find_skills(raw_text, FLATTENED_SKILL_POOL, threshold=78)
    specialization, scores = get_specialization(skills, education, experience)

    summary_parts = []
    if name != "Not Found":
        summary_parts.append(f"A profile for {name} has been successfully parsed. The system uses an advanced NLP pipeline.")
    if education:
        summary_parts.append(f"Education highlights include: {education[0] + (', ...' if len(education) > 1 else '')}.")
    if experience:
        summary_parts.append(f"The experience section details roles such as '{experience[0].split('(')[0].strip()}' and mentions key projects.")
    if skills:
        summary_parts.append(
            f"The candidate's core skills are in **{specialization}**, with a strong command of: {', '.join(skills[:6]) + '...' if len(skills) > 6 else ''}."
        )
    summary_parts.append(f"Overall, the resume suggests a strong fit for a **{specialization}** role.")

    summary = "\n\n".join(summary_parts)

    return {
        "Name": name,
        "Emails": emails,
        "Phones": phones,
        "Education": education,
        "Experience": experience,
        "Skills": skills,
        "Specialization": specialization,
        "SpecializationScores": scores,
        "Summary": summary,
    }

# -------- Route --------
@app.route('/analyze', methods=['POST'])
def analyze_resume():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded."}), 400
    file = request.files['file']
    ext = os.path.splitext(file.filename)[1].lower()
    temp_path = "temp" + ext
    file.save(temp_path)
    try:
        if ext == ".pdf":
            text = extract_text_from_pdf(temp_path)
        elif ext == ".docx":
            text = extract_text_from_docx(temp_path)
        elif ext == ".txt":
            with open(temp_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        else:
            os.remove(temp_path)
            return jsonify({"error": "Unsupported file type"}), 400
        result = parse_resume_text(text)
        os.remove(temp_path)
        return jsonify(result)
    except Exception as e:
        os.remove(temp_path)
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    app.run(debug=True, port=5001)
